"""Authored fixtures are generated locally; no external dataset is downloaded."""
from pathlib import Path
import pytest


@pytest.fixture(scope='session')
def fixture_dir(tmp_path_factory):
    import fitz
    from docx import Document
    from openpyxl import Workbook
    root = tmp_path_factory.mktemp('authored-documents')
    for name, temperature, status in [('spec_r2.pdf', 94, 'current'),
                                      ('spec_r1_WITHDRAWN.pdf', 114, 'withdrawn')]:
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((54, 70), f'Product: NX-57\nMaximum junction temperature: {temperature}\nStatus: {status}')
            document.save(root / name)
    with fitz.open() as document:
        document.new_page().insert_text((54, 70), 'Product: NX-57\nUnit price: 1')
        document.save(root / 'supplier_ENCRYPTED.pdf', encryption=fitz.PDF_ENCRYPT_AES_256,
                      owner_pw='authored-owner-password', user_pw='authored-user-password')
    document = Document()
    document.add_paragraph('NX-57 product roadmap')
    table = document.add_table(rows=0, cols=3)
    for row in [['Product', 'Milestone', 'Quarter'], ['NX-57', 'Customer sampling', 'Q3 FY29'],
                ['NX-58', 'Customer sampling', 'Q1 FY30']]:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    document.save(root / 'roadmap.docx')
    workbook = Workbook()
    parts = workbook.active
    parts.title = 'Parts'
    for row in [['Product', 'Description', 'Part number', 'Lead time days'],
                ['NX-57', 'field replaceable fan assembly', 'ORBIT-FAN-7781-B', 12],
                ['NX-57', 'power module', 'ORBIT-PSU-0082-A', 9],
                ['NX-58', 'field replaceable fan assembly', 'ORBIT-FAN-1228-A', 4]]:
        parts.append(row)
    prices = workbook.create_sheet('Prices')
    for row in [['Product', 'Volume', 'Unit price'], ['NX-57', 1000, 129], ['NX-57', 10000, 94]]:
        prices.append(row)
    workbook.save(root / 'parts.xlsx')
    workbook.close()
    return root
