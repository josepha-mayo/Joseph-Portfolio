"""DOCX merged-cell scope must survive literal OOXML parsing."""
from docx import Document
from von_rag.parsers import parse_docx

def make_vertical(path):
    d=Document(); table=d.add_table(rows=4,cols=3)
    for j,value in enumerate(['Product','Property','Value']):table.cell(0,j).text=value
    table.cell(1,0).text='TQ-40'
    table.cell(1,1).text='Maximum junction temperature';table.cell(1,2).text='94'
    table.cell(2,1).text='Board revision';table.cell(2,2).text='REV-C2'
    table.cell(1,0).merge(table.cell(2,0))
    table.cell(3,0).text='TQ-60';table.cell(3,1).text='Board revision';table.cell(3,2).text='REV-D4'
    d.save(path)

def test_vertical_merge_carries_product_scope_then_stops(tmp_path):
    path=tmp_path/'merged.docx';make_vertical(path)
    rows=[c for c in parse_docx(path,'planning/merged.docx') if c.kind=='row']
    assert len(rows)==3
    assert rows[0].fields=={'Product':'TQ-40','Property':'Maximum junction temperature','Value':'94'}
    assert rows[1].fields=={'Product':'TQ-40','Property':'Board revision','Value':'REV-C2'}
    assert rows[2].fields=={'Product':'TQ-60','Property':'Board revision','Value':'REV-D4'}

def test_merged_continuation_is_retrievable_by_product(tmp_path):
    path=tmp_path/'merged.docx';make_vertical(path)
    rows=parse_docx(path,'planning/merged.docx')
    revision=next(c for c in rows if c.fields.get('Value')=='REV-C2')
    assert revision.fields['Product']=='TQ-40'
    assert 'Product: TQ-40' in revision.text

def test_horizontal_grid_span_keeps_later_columns_aligned(tmp_path):
    d=Document(); table=d.add_table(rows=3,cols=3)
    title=table.cell(0,0).merge(table.cell(0,2));title.text='Roadmap FY31'
    for j,value in enumerate(['Product','Property','Value']):table.cell(1,j).text=value
    table.cell(2,0).text='DN-315';table.cell(2,1).text='Customer sampling quarter';table.cell(2,2).text='Q3 FY31'
    path=tmp_path/'horizontal.docx';d.save(path)
    rows=[c for c in parse_docx(path,'planning/horizontal.docx') if c.kind=='row']
    assert len(rows)==1
    assert rows[0].fields['Product']=='DN-315'
    assert rows[0].fields['Property']=='Customer sampling quarter'
    assert rows[0].fields['Value']=='Q3 FY31'
    assert 'Roadmap FY31' in rows[0].context
