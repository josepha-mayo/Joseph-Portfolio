"""Five-slide updated pitch in editable PowerPoint and accessible HTML."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
import html
slides=[
 ('The event vanished.\nThe ticket did not.','FORKLINE / DELIVERY LAB','Test the gap between queue approval and outside work. A reproducible local integration for Web3 ticketing, membership and fulfillment developers.'),
 ('Approval is not a receipt.','01 / THE FAILURE','An event can be eligible at enqueue and orphaned at dispatch. A receiver can commit while its acknowledgement is lost. Neither rolling back an index nor blindly retrying resolves both cases.'),
 ('A real worker. A separate receiver.','02 / THE BUILD','Compiled Solidity fixture and pinned EVM observations → shared eligibility engine → SQLite outbox → loopback HTTP receiver → separate ticket ledger. Restart the worker and reconcile its receipt.'),
 ('Reproduce the failure.\nKeep the outside effect.','03 / THE DIFFERENCE','Same waiting rule, same order: the enqueue-only consumer issues one orphaned fixture ticket; Forkline holds the command. If a ticket already exists, retain its receipt, latch the incident and stop further guarded delivery.'),
 ('An integration test you can rerun.','04 / EVIDENCE & NEXT STEP','Executed database and HTTP checks include dropped acknowledgements, true child-process death, restart, receipt lookup and duplicate commands. No real assets or field savings. Next: maintained live-node adapter and an authorized developer pilot.')]
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
for i,(title,tag,body)in enumerate(slides):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('0B1017')
 for txt,y,h,size,color in [(tag,.65,.5,13,'83E8CE'),(title,1.45,2.0,36,'EFF5F9'),(body,3.8,2.2,22,'A6B5C3'),(f'FORKLINE  /  {i+1:02d} OF 05    ·    LOCAL TEST FIXTURES    ·    MIT    ·    AI-ASSISTED ORIGINAL WORK',6.8,.3,10,'A6B5C3')]:
  box=s.shapes.add_textbox(Inches(.85),Inches(y),Inches(11.6),Inches(h));tf=box.text_frame;tf.word_wrap=True;tf.text=txt
  for p in tf.paragraphs:p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color)
prs.save('public/pitch.pptx')
css='body{margin:0;background:#0b1017;color:#eff5f9;font:18px/1.5 system-ui}section{min-height:100vh;box-sizing:border-box;padding:9vh 9vw;display:flex;flex-direction:column;justify-content:center;border-bottom:1px solid #293848}h1{font-size:clamp(32px,5vw,64px);line-height:1.1;max-width:1050px;margin:24px 0 32px}p{max-width:1000px;color:#a6b5c3;font-size:clamp(18px,2vw,26px)}small,a{color:#83e8ce}footer{font-size:13px;margin-top:35px}'
page='<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Forkline Delivery Lab | Pitch</title><style>'+css+'</style>'
for i,(title,tag,body)in enumerate(slides):page+=f'<section><small>{html.escape(tag)}</small><h1>{html.escape(title).replace(chr(10),"<br>")}</h1><p>{html.escape(body)}</p><footer>{i+1:02d} / 05 · <a href="delivery.html">Delivery Lab</a> · <a href="pitch.pptx">PowerPoint</a></footer></section>'
Path('public/pitch.html').write_text(page+'</html>')
